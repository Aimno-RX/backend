import sys, os
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, '.')
from docx import Document
from docx.oxml.ns import qn

file_path = r'E:\xwechat_files\wxid_kngvfm4spkj422_1a7f\msg\file\2026-06\《颈椎胸椎功能强化训练》已整理完毕.docx'
doc = Document(file_path)

# Find major section boundaries by looking at pattern changes
# The document follows: 练习编号 → 起始姿势 → [动作图片] → 动作要领 → 重复次数 → 注意事项
# We need to find the major section headers (Day numbers, topic headers)

# Look for text that indicates major section breaks
major_sections = []
prev_bold = ""
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    runs = para.runs
    is_bold = any(r.bold for r in runs if r.text.strip()) if runs else False
    
    if is_bold:
        # Detect section-level headers (not exercise-level)
        section_markers = [
            '颈椎', '胸椎', '训练', '天练习', '平躺', '坐姿', '站立',
            '肩胛', '科普', '小常识', '小贴士', '马车夫', '头部',
            'B ', '目标', '为什么', '安全', '进阶', '综合',
            '腹肌', '背肌', '拉伸', '放松', '强化'
        ]
        for marker in section_markers:
            if marker in text and not text.startswith('练习') and not text.startswith('起始') and not text.startswith('动作要领') and not text.startswith('重复') and not text.startswith('注意事项'):
                major_sections.append((i, text[:80]))
                break

print("=== MAJOR SECTIONS ===\n")
current = ""
for idx, (para_i, text) in enumerate(major_sections):
    if text != current:
        print(f"p{para_i:04d}: {text}")
        current = text

# Now map images to their nearby exercise names
print("\n\n=== IMAGE-TO-EXERCISE MAPPING (sample) ===\n")

# For each paragraph with an image, find the nearest bold text before it that contains "练习" or section name
img_exercises = []
for i, para in enumerate(doc.paragraphs):
    blips = para._element.findall('.//' + qn('a:blip'))
    if not blips:
        continue
    
    # Search backwards for the exercise name
    exercise_name = "未知"
    section_name = "未知"
    for j in range(i-1, max(0, i-50), -1):
        prev_text = doc.paragraphs[j].text.strip()
        prev_runs = doc.paragraphs[j].runs
        prev_bold = any(r.bold for r in prev_runs if r.text.strip()) if prev_runs else False
        if prev_bold and prev_text:
            if '练习' in prev_text or '起始' in prev_text:
                exercise_name = prev_text
                break
            if prev_text and exercise_name == "未知":
                section_name = prev_text
                break
    
    img_exercises.append((i, len(blips), exercise_name, section_name))

# Group by exercise name and count
from collections import Counter
exercise_counts = Counter()
for para_i, count, exercise, section in img_exercises:
    exercise_counts[exercise] += count

print("Images per exercise type:")
for name, count in exercise_counts.most_common(30):
    print(f"  [{count} imgs] {name}")

# Show sample images with their context
print("\n\n=== DETAILED IMAGE CONTEXT (first 20) ===\n")
for para_i, count, exercise, section in img_exercises[:20]:
    print(f"p{para_i:04d}: [{count} img] near '{exercise}' | section: '{section}'")

print("\n\n=== DETAILED IMAGE CONTEXT (last 20) ===\n")
for para_i, count, exercise, section in img_exercises[-20:]:
    print(f"p{para_i:04d}: [{count} img] near '{exercise}' | section: '{section}'")
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
os.environ['IMAGE_STORAGE_DIR'] = './test_images'

from dotenv import load_dotenv
load_dotenv(override=True)

from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage
import base64

# Config
vision_config = os.getenv('LLM_MODEL_CONFIG_qwen_vl', '')
model_name, api_endpoint, api_key = vision_config.split(',')
vision_llm = ChatOpenAI(api_key=api_key, base_url=api_endpoint, model=model_name, temperature=0, max_tokens=2048, request_timeout=120)

# Extract images from docx
from docx import Document
from lxml import etree
from PIL import Image
import io

file_path = r'E:\xwechat_files\wxid_kngvfm4spkj422_1a7f\msg\file\2026-06\《颈椎胸椎功能强化训练》已整理完毕.docx'
MAX_IMAGE_SIZE = 2048

def resize_if_needed(image_bytes):
    img = Image.open(io.BytesIO(image_bytes))
    w, h = img.size
    if w <= MAX_IMAGE_SIZE and h <= MAX_IMAGE_SIZE:
        return image_bytes, img.format or 'PNG'
    ratio = min(MAX_IMAGE_SIZE / w, MAX_IMAGE_SIZE / h)
    new_size = (int(w * ratio), int(h * ratio))
    img = img.resize(new_size, Image.LANCZOS)
    buf = io.BytesIO()
    img_format = img.format or 'PNG'
    img.save(buf, format=img_format)
    return buf.getvalue(), img_format

doc = Document(file_path)
nsmap = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}

image_part_map = {}
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_part_map[rel.rId] = {
            "blob": rel.target_part.blob,
            "ext": rel.target_part.content_type.split("/")[-1],
        }

# Extract up to 5 images
images = []
for para_idx, paragraph in enumerate(doc.paragraphs):
    if len(images) >= 5:
        break
    blips = paragraph._element.findall('.//' + etree.QName(nsmap['a'], 'blip').text, nsmap)
    for blip in blips:
        if len(images) >= 5:
            break
        embed = blip.get(etree.QName(nsmap['r'], 'embed').text)
        if embed and embed in image_part_map:
            img_info = image_part_map[embed]
            image_bytes = img_info["blob"]
            ext = img_info["ext"]
            if ext == 'x-emf':
                print(f"  Skipping EMF image at p{para_idx} (unsupported format)")
                continue
            if len(image_bytes) < 1024:
                continue
            resized_bytes, fmt = resize_if_needed(image_bytes)
            images.append({
                "paragraph_index": para_idx,
                "image_bytes": resized_bytes,
                "extension": ext,
                "index": len(images),
            })

print(f"Extracted {len(images)} testable images\n")

# Describe each image
for img in images:
    para_idx = img['paragraph_index']
    ext = img['extension']
    mime = f"image/{ext}" if ext != 'jpg' else 'image/jpeg'
    b64 = base64.b64encode(img['image_bytes']).decode('utf-8')
    
    # Get surrounding text context
    context_parags = []
    for j in range(max(0, para_idx - 3), min(len(doc.paragraphs), para_idx + 3)):
        txt = doc.paragraphs[j].text.strip()
        if txt:
            context_parags.append(txt[:100])
    context = "\n".join(context_parags[:5])
    
    prompt = f"""This image appears in a rehabilitation / physical therapy training document near paragraph {para_idx}.

The surrounding text context is:
{context}

Please describe this image as a rehabilitation exercise guide. Focus on extracting:
1. **训练动作 (Exercise Name)**: What specific exercise or movement is being demonstrated?
2. **姿势 (Posture/Body Position)**: What is the starting position? Sitting, standing, lying down?
3. **动作描述 (Movement Description)**: Describe the movement trajectory step by step.
4. **训练参数 (Training Parameters)**: Repetitions, sets, hold time.
5. **针对部位 (Target Body Parts)**: Which muscles, joints are being targeted?
6. **注意事项 (Precautions)**: Any safety cues or warnings.

Provide the description in Chinese."""

    print(f"=== Image {img['index']+1} at paragraph {para_idx} ({len(img['image_bytes'])//1024}KB, {ext}) ===")
    
    try:
        message = HumanMessage(content=[
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
        ])
        response = vision_llm.invoke([message])
        print(f"Description:\n{response.content}\n")
    except Exception as e:
        print(f"ERROR: {e}\n")

print("=== TEST COMPLETE ===")